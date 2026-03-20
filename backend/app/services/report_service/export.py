from io import BytesIO
from sqlalchemy.orm import Session
from sqlalchemy import func

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.models.time_entry import TimeEntry
from app.models.project import Project
from app.models.user import User
from app.models.task import Task
from .query import apply_filters

def generate_xlsx_export(db: Session, **filters) -> bytes:
    # 1. Nueva consulta granular: agrupamos por Proyecto, Fecha, Usuario y Tarea
    query = db.query(
        Project.code.label("project_code"),
        Project.name.label("project_name"),
        Project.distance_from_workshop.label("project_distance"),
        TimeEntry.date.label("date"),
        TimeEntry.is_holiday.label("is_holiday"),
        User.employee_code.label("employee_code"),
        User.name.label("employee_name"),
        Task.code.label("task_code"),
        Task.name.label("task_name"),
        TimeEntry.vehicle_type.label("vehicle_type"),
        TimeEntry.meals.label("meals"),
        TimeEntry.distance_origin.label("distance_origin"),
        TimeEntry.trip_type.label("trip_type"),
        func.coalesce(func.sum(TimeEntry.hours + TimeEntry.overtime_hours), 0).label("total_hours"),
        func.coalesce(func.sum(TimeEntry.travel_time), 0).label("travel_time")
    ).join(Project, TimeEntry.project_id == Project.id)\
     .join(Task, TimeEntry.task_id == Task.id)\
     .join(User, TimeEntry.user_id == User.id)

    query = apply_filters(query, **filters)
    
    # Orden lógico: Proyecto -> Fecha -> Usuario -> Tarea
    entries = query.group_by(
        Project.code,
        Project.name,
        Project.distance_from_workshop,
        TimeEntry.date,
        TimeEntry.is_holiday,
        User.employee_code,
        User.name,
        Task.code,
        Task.name,
        TimeEntry.vehicle_type,
        TimeEntry.meals,
        TimeEntry.distance_origin,
        TimeEntry.trip_type
    ).order_by(Project.code, TimeEntry.date, User.employee_code, Task.code).all()

    # Reestructuramos los datos para enviarlos al generador de la hoja:
    projects_dict = {}
    for entry in entries:
        p_code = entry.project_code
        if p_code not in projects_dict:
            projects_dict[p_code] = {
                "code": p_code,
                "name": entry.project_name,
                "rows": []
            }
        
        projects_dict[p_code]["rows"].append({
            "date": entry.date,
            "is_holiday": entry.is_holiday,
            "emp_code": entry.employee_code,
            "emp_name": entry.employee_name,
            "task_code": entry.task_code,
            "task_name": entry.task_name,
            "hours": float(entry.total_hours),
            "travel_time": float(entry.travel_time or 0),
            "vehicle": entry.vehicle_type,
            "meals": entry.meals,
            "origin": entry.distance_origin,
            "trip_type": entry.trip_type,
            "project_distance": entry.project_distance
        })

    projects_list = list(projects_dict.values())
    return _build_openpyxl_wb(projects_list)

def _build_openpyxl_wb(projects: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Informe Horas"

    # ── Estilos ──
    header_font = Font(name="Arial", bold=True, size=10)
    header_fill = PatternFill("solid", fgColor="FFFF00")  # amarillo
    header_alignment = Alignment(horizontal="center", vertical="center")

    data_font = Font(name="Arial", size=10)
    data_alignment_left = Alignment(horizontal="left", vertical="center")
    data_alignment_center = Alignment(horizontal="center", vertical="center")
    data_alignment_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # ── Anchos de columna ──
    col_widths = {
        "A": 12,   # Fecha
        "B": 10,   # Festivo
        "C": 14,   # Cód. empleado
        "D": 25,   # Nombre
        "E": 12,   # Cód. artículo
        "F": 25,   # Desc. artículo
        "G": 10,   # Horas
        "H": 12,   # Desplaz. (h)
        "I": 10,   # Precio
        "J": 15,   # Transporte
        "K": 15,   # Origen
        "L": 10,   # KMs
        "M": 10,   # Dieta
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    HEADERS = ["Fecha", "Festivo", "Cód. empleado", "Nombre", "Cód. artículo", "Desc. artículo", 
               "Horas", "Desplaz. (h)", "Precio", "Transporte", "Origen", "KMs", "Dieta"]

    # Style for holiday rows
    holiday_fill = PatternFill("solid", fgColor="FFC7CE")  # light red

    current_row = 1

    for proj_idx, project in enumerate(projects):
        if proj_idx > 0:
            current_row += 1  # fila en blanco entre proyectos

        # Título del proyecto
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(HEADERS))
        title_cell = ws.cell(row=current_row, column=1)
        title_cell.value = f"Proyecto: {project['code']} - {project['name']}"
        title_cell.font = Font(name="Arial", bold=True, size=11, color="333333")
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        title_cell.fill = PatternFill("solid", fgColor="D9E1F2")  # azul claro
        for col in range(1, len(HEADERS) + 1):
            ws.cell(row=current_row, column=col).border = thin_border
        current_row += 1

        # ── Cabecera amarilla ──
        for col_idx, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        current_row += 1

        # ── Filas de datos agrupados por día y trabajador ──
        first_data_row = current_row
        for row_info in project["rows"]:
            
            # Castings estéticos
            emp_code = str(row_info["emp_code"])
            task_code = str(row_info["task_code"])
            is_holiday = row_info.get("is_holiday", False)
            
            # Mapeo de términos para mejor lectura
            transporte = "Particular" if row_info["vehicle"] in ("personal", "particular") else ("Empresa" if row_info["vehicle"] in ("company", "empresa") else "-")
            
            tipo_trayecto = ""
            if row_info["trip_type"] == "to": tipo_trayecto = "(IDA)"
            elif row_info["trip_type"] == "from": tipo_trayecto = "(VUELTA)"
            elif row_info["trip_type"] == "round": tipo_trayecto = "(IDA+VTA)"
            
            origin_base = "NAVE" if row_info["origin"] == "workshop" else (str(row_info["origin"]) if row_info["origin"] else "-")
            origen = f"{origin_base} {tipo_trayecto}".strip()
            
            dieta = "SÍ" if row_info["meals"] else ("NO" if row_info["meals"] is False else "-")
            
            # Cálculo de KM: Solo si el coche es PARTICULAR
            kms = 0.0
            if row_info["vehicle"] in ("personal", "particular"):
                # Multiplicador: IDA=1, VUELTA=1, AMBOS=2
                multiplier = 2.0 if row_info["trip_type"] == "round" else 1.0
                dist_base = float(row_info["project_distance"] or 0.0)
                kms = dist_base * multiplier
            
            row_data = [
                row_info["date"].strftime("%d-%m-%Y"),                 # Fecha
                "SÍ" if is_holiday else "",                            # Festivo
                int(emp_code) if emp_code.isdigit() else emp_code,     # Cód. empleado
                row_info["emp_name"],                                  # Nombre operario
                int(task_code) if task_code.isdigit() else task_code,  # Cód. artículo (tarea)
                row_info["task_name"],                                 # Desc. artículo
                row_info["hours"],                                     # Horas
                row_info.get("travel_time", 0),                        # Desplaz. (h)
                0.00,                                                  # Precio (placeholder)
                transporte,                                            # Transporte
                origen,                                                # Origen
                kms,                                                   # KMs
                dieta,                                                 # Dieta
            ]

            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.value = value
                cell.font = data_font
                cell.border = thin_border

                if is_holiday:
                    cell.fill = holiday_fill

                if col_idx == 1:  # Fecha
                    cell.alignment = data_alignment_center
                elif col_idx == 2:  # Festivo
                    cell.alignment = data_alignment_center
                elif col_idx in (3, 5):  # códigos numéricos
                    cell.alignment = data_alignment_center
                elif col_idx in (4, 6, 10, 11, 13):  # textos y etiquetas
                    cell.alignment = data_alignment_left
                elif col_idx in (7, 8, 9, 12):  # Horas / Desplaz. / Precio / KMs
                    cell.alignment = data_alignment_right
                    cell.number_format = '#,##0.00'

            current_row += 1
            
        if first_data_row == current_row:
            last_data_row = first_data_row
            ws.cell(row=first_data_row, column=1).value = "Sin registros"
            current_row += 1
        else:
            last_data_row = current_row - 1

        # ── Fila de TOTAL por proyecto ──
        total_font = Font(name="Arial", bold=True, size=10)
        total_fill = PatternFill("solid", fgColor="E2EFDA")  # verde claro

        ws.cell(row=current_row, column=1).value = ""
        ws.cell(row=current_row, column=1).border = thin_border
        ws.cell(row=current_row, column=1).fill = total_fill

        for col in range(2, 6):
            c = ws.cell(row=current_row, column=col)
            c.value = ""
            c.border = thin_border
            c.fill = total_fill

        ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=5)
        label_cell = ws.cell(row=current_row, column=2)
        label_cell.value = f"TOTAL {project['code']}"
        label_cell.font = total_font
        label_cell.alignment = Alignment(horizontal="right", vertical="center")
        label_cell.fill = total_fill

        # Fórmula SUM para Horas (col G)
        sum_cell = ws.cell(row=current_row, column=7)
        sum_cell.value = f"=SUM(G{first_data_row}:G{last_data_row})"
        sum_cell.font = total_font
        sum_cell.alignment = data_alignment_right
        sum_cell.number_format = '#,##0.00'
        sum_cell.border = thin_border
        sum_cell.fill = total_fill

        # Fórmula SUM para Desplaz. (col H)
        travel_cell = ws.cell(row=current_row, column=8)
        travel_cell.value = f"=SUM(H{first_data_row}:H{last_data_row})"
        travel_cell.font = total_font
        travel_cell.alignment = data_alignment_right
        travel_cell.number_format = '#,##0.00'
        travel_cell.border = thin_border
        travel_cell.fill = total_fill

        # Fórmula SUM para Precio (col I)
        price_cell = ws.cell(row=current_row, column=9)
        price_cell.value = f"=SUM(I{first_data_row}:I{last_data_row})"
        price_cell.font = total_font
        price_cell.alignment = data_alignment_right
        price_cell.number_format = '#,##0.00'
        price_cell.border = thin_border
        price_cell.fill = total_fill
        
        # Columna de totales para KMs (col L)
        km_total_cell = ws.cell(row=current_row, column=12)
        km_total_cell.value = f"=SUM(L{first_data_row}:L{last_data_row})"
        km_total_cell.font = total_font
        km_total_cell.alignment = data_alignment_right
        km_total_cell.number_format = '#,##0.00'
        km_total_cell.border = thin_border
        km_total_cell.fill = total_fill

        current_row += 1

    # ── Configuración de impresión ──
    ws.sheet_properties.pageSetUpPr = None
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.oddHeader.center.text = "TimeFlow - Informe Diario de Horas por Proyecto"
    ws.oddFooter.right.text = "Página &P de &N"

    # Save to BytesIO stream
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
